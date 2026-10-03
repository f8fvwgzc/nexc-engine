"""safe_path confinement, workspace quotas, make_docx output and strict schemas."""

from __future__ import annotations

import io
import os
from pathlib import Path

import docx
import pytest

from nexc_runtime.tools import PathError, Workspace, build_toolset, render_docx, safe_path
from nexc_runtime.tools.docx_tool import MakeDocxArgs


@pytest.mark.parametrize(
    "bad",
    [
        "",
        "   ",
        "/etc/passwd",
        "../outside.txt",
        "a/../../outside.txt",
        "a/b/../../../x",
        "C:/Windows/win.ini",
        "dir\\file.txt",
        "nul\x00byte",
        "x" * 600,
    ],
)
def test_safe_path_rejects_escapes(tmp_path: Path, bad: str) -> None:
    with pytest.raises(PathError):
        safe_path(tmp_path, bad)


def test_safe_path_rejects_symlink_escape(tmp_path: Path) -> None:
    root = tmp_path / "ws"
    root.mkdir()
    outside = tmp_path / "secret.txt"
    outside.write_text("secret")
    os.symlink(outside, root / "link.txt")
    os.symlink(tmp_path, root / "linkdir")
    with pytest.raises(PathError):
        safe_path(root, "link.txt")
    with pytest.raises(PathError):
        safe_path(root, "linkdir/secret.txt")


def test_safe_path_accepts_nested_relative_paths(tmp_path: Path) -> None:
    resolved = safe_path(tmp_path, "src/pkg/./module.py")
    assert resolved == (tmp_path / "src" / "pkg" / "module.py").resolve()


def test_workspace_enforces_quotas(tmp_path: Path) -> None:
    ws = Workspace(tmp_path, max_file_bytes=10, max_total_bytes=15)
    ws.write_bytes("a.txt", b"0123456789")
    with pytest.raises(PathError):
        ws.write_bytes("b.txt", b"0123456789")  # total quota
    with pytest.raises(PathError):
        ws.write_bytes("c.txt", b"x" * 11)  # per-file limit
    ws.write_bytes("a.txt", b"short")  # overwriting frees the old size


def test_make_docx_produces_a_valid_document() -> None:
    args = MakeDocxArgs.model_validate(
        {
            "filename": "report.docx",
            "title": "Market Analysis",
            "sections": [
                {
                    "heading": "Summary",
                    "paragraphs": ["Revenue grew 12%.", "Margins held."],
                    "bullets": ["Expand to EU", "Hire two engineers"],
                    "table_rows": [["Quarter", "Revenue"], ["Q1", "1.2M"], ["Q2", "1.4M", "extra"]],
                },
                {"heading": "Risks", "paragraphs": [], "bullets": ["FX"], "table_rows": []},
            ],
        }
    )
    document = docx.Document(io.BytesIO(render_docx(args)))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "Market Analysis"
    assert "Summary" in texts and "Risks" in texts
    assert "Revenue grew 12%." in texts
    bullets = [
        p.text for p in document.paragraphs if p.style is not None and p.style.name == "List Bullet"
    ]
    assert bullets == ["Expand to EU", "Hire two engineers", "FX"]
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(0, 0).text == "Quarter"
    assert table.cell(1, 2).text == ""  # ragged rows are padded
    assert document.core_properties.title == "Market Analysis"


def test_tool_schemas_are_strict() -> None:
    tools = build_toolset(can_spawn=True, allow_code_exec=True)

    def check(node: object) -> None:
        if isinstance(node, dict):
            assert "$ref" not in node
            assert "title" not in node
            if node.get("type") == "object":
                assert node["additionalProperties"] is False
            for key, value in node.items():
                if key == "properties":
                    for prop in value.values():
                        check(prop)
                else:
                    check(value)
        elif isinstance(node, list):
            for item in node:
                check(item)

    for tool in tools.values():
        schema = tool.spec().input_schema
        check(schema)
        assert set(schema["required"]) == set(schema["properties"])


def test_run_python_is_disabled_by_default() -> None:
    assert "run_python" not in build_toolset(can_spawn=True, allow_code_exec=False)
    assert "run_python" in build_toolset(can_spawn=True, allow_code_exec=True)


def test_safe_path_dot_is_the_workspace_root(tmp_path: Path) -> None:
    assert safe_path(tmp_path, ".") == tmp_path.resolve()
    assert safe_path(tmp_path, "./") == tmp_path.resolve()
