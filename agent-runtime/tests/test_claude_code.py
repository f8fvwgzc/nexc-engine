"""The `claude_code` provider drives the tool loop through the Claude Code CLI protocol."""

from __future__ import annotations

import asyncio
import base64
import io
import json
import os
import stat
from collections.abc import Callable
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from nexc_runtime.agents import agent
from nexc_runtime.api.app import create_app
from nexc_runtime.config import Settings
from nexc_runtime.llm.claude_code import ClaudeCodeProvider, describe_failure

from .conftest import AUTH, execute_body, parse_ndjson

# A stand-in for `claude -p --output-format json --json-schema ...`: first turn asks for
# make_docx, the turn after the tool result finishes. It also records its argv and environment.
FAKE_CLI = """#!/usr/bin/env python3
import json, os, sys
prompt = sys.stdin.read()
with open(os.environ["HOME"] + "/calls.jsonl", "a") as log:
    log.write(json.dumps({"argv": sys.argv[1:], "env": sorted(os.environ)}) + "\\n")
if "### Action results" in prompt:
    out = {"text": "Report written to report.docx.", "actions": []}
else:
    out = {"text": "Drafting.", "actions": [{"action": "make_docx", "input": {
        "filename": "report.docx", "title": "Q3 report",
        "sections": [{"heading": "Summary", "paragraphs": ["Revenue grew."],
                      "bullets": [], "table_rows": []}]}}]}
print(json.dumps({"type": "result", "subtype": "success", "is_error": False, "result": "",
    "stop_reason": "end_turn", "structured_output": out,
    "usage": {"input_tokens": 5, "cache_read_input_tokens": 5, "output_tokens": 7}}))
"""


def _fake_cli(tmp_path: Path) -> Path:
    path = tmp_path / "claude"
    path.write_text(FAKE_CLI)
    path.chmod(path.stat().st_mode | stat.S_IEXEC)
    return path


def test_claude_code_runs_tools_and_produces_artifacts(
    make_settings: Callable[..., Settings], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("NEXC_RUNTIME_TOKEN", "must-not-leak-into-the-cli")
    settings = make_settings(runtime_claude_bin=str(_fake_cli(tmp_path)))
    with TestClient(create_app(settings)) as client:
        body = execute_body(llm={"provider": "claude_code", "api_key": None, "model": "haiku"})
        response = client.post("/v1/execute", json=body, headers=AUTH)

    events = parse_ndjson(response.content)
    assert events[-1]["type"] == "result", events[-1]
    assert events[-1]["output"] == "Report written to report.docx."
    assert events[-1]["tokens_in"] == 20 and events[-1]["tokens_out"] == 14
    (artifact,) = [e for e in events if e["type"] == "artifact"]
    document = docx.Document(io.BytesIO(base64.b64decode(artifact["content_b64"])))
    assert "Revenue grew." in [p.text for p in document.paragraphs]

    calls = [json.loads(line) for line in (tmp_path / "calls.jsonl").read_text().splitlines()]
    assert len(calls) == 2
    argv = calls[0]["argv"]
    assert argv[argv.index("--model") + 1] == "haiku"
    assert argv[argv.index("--tools") + 1] == "", "built-in CLI tools are disabled"
    schema = json.loads(argv[argv.index("--json-schema") + 1])
    assert "make_docx" in schema["properties"]["actions"]["items"]["properties"]["action"]["enum"]
    assert "NEXC_RUNTIME_TOKEN" not in calls[0]["env"], "secrets never reach the CLI"


def test_missing_cli_is_a_clear_non_retryable_error(
    make_settings: Callable[..., Settings],
) -> None:
    settings = make_settings(runtime_claude_bin="/nonexistent/claude")
    with TestClient(create_app(settings)) as client:
        body = execute_body(llm={"provider": "claude_code", "api_key": None})
        events = parse_ndjson(client.post("/v1/execute", json=body, headers=AUTH).content)
    assert events[-1]["type"] == "error"
    assert "install Claude Code" in events[-1]["message"]
    assert events[-1]["retryable"] is False


def test_slow_turns_emit_heartbeats(
    make_settings: Callable[..., Settings], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A provider that answers only at the end of a turn must not leave the stream silent."""
    monkeypatch.setattr(agent, "HEARTBEAT_EVERY_S", 0.05)
    slow = tmp_path / "claude"
    slow.write_text(
        FAKE_CLI.replace(
            "prompt = sys.stdin.read()", "prompt = sys.stdin.read()\nimport time; time.sleep(0.3)"
        )
    )
    slow.chmod(slow.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("HOME", str(tmp_path))
    settings = make_settings(runtime_claude_bin=str(slow))
    with TestClient(create_app(settings)) as client:
        body = execute_body(llm={"provider": "claude_code", "api_key": None})
        events = parse_ndjson(client.post("/v1/execute", json=body, headers=AUTH).content)
    assert events[-1]["type"] == "result"
    beats = [e for e in events if e["type"] == "log" and "still working" in e["message"]]
    assert len(beats) >= 2, events


def test_text_only_document_answers_still_deliver_a_docx(
    make_settings: Callable[..., Settings], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A document node whose agent answered in Markdown without make_docx gets a converted .docx."""
    answer = "\n".join(
        [
            "# Solid-state batteries 2026",
            "",
            "Intro paragraph.",
            "",
            "## Findings",
            "- **Energy** density up",
            "1. Second point",
            "",
            "| Maker | Year |",
            "|---|---|",
            "| Toyota | 2027 |",
        ]
    )
    cli = tmp_path / "claude"
    cli.write_text(
        FAKE_CLI.replace('if "### Action results" in prompt:', "if True:").replace(
            '"Report written to report.docx."', json.dumps(answer)
        )
    )
    cli.chmod(cli.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("HOME", str(tmp_path))
    settings = make_settings(runtime_claude_bin=str(cli))
    with TestClient(create_app(settings)) as client:
        body = execute_body(
            llm={"provider": "claude_code", "api_key": None},
            task={"title": "Final report", "kind": "output"},
        )
        events = parse_ndjson(client.post("/v1/execute", json=body, headers=AUTH).content)
    (artifact,) = [e for e in events if e["type"] == "artifact"]
    assert artifact["path"] == "final-report.docx"
    document = docx.Document(io.BytesIO(base64.b64decode(artifact["content_b64"])))
    text = [p.text for p in document.paragraphs]
    assert text[0] == "Solid-state batteries 2026"
    assert "Findings" in text and "Energy density up" in text and "Second point" in text
    assert document.tables[0].cell(1, 0).text == "Toyota"


def test_cli_failures_report_subtype_and_errors() -> None:
    message, retryable = describe_failure(
        {"subtype": "error_max_turns", "is_error": True, "errors": ["Reached maximum number"]}
    )
    assert message == "claude CLI: error_max_turns: Reached maximum number"
    assert retryable
    message, retryable = describe_failure(
        {"subtype": "success", "is_error": True, "result": "Invalid API key · Please run /login"}
    )
    assert "Invalid API key" in message
    assert not retryable


FAILING_SCHEMA_CLI = """#!/usr/bin/env python3
import json, sys
sys.stdin.read()
if "--json-schema" in sys.argv:
    print(json.dumps({"type": "result", "subtype": "error_max_structured_output_retries",
        "is_error": True, "errors": ["Failed to provide valid structured output"]}))
else:
    print(json.dumps({"type": "result", "subtype": "success", "is_error": False,
        "result": "# Sources\\n\\nA long plain answer.", "stop_reason": "end_turn",
        "usage": {"input_tokens": 3, "output_tokens": 4}}))
"""


def test_structured_output_failure_falls_back_to_a_plain_answer(
    make_settings: Callable[..., Settings], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    cli = tmp_path / "claude"
    cli.write_text(FAILING_SCHEMA_CLI)
    cli.chmod(cli.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("HOME", str(tmp_path))
    settings = make_settings(runtime_claude_bin=str(cli))
    with TestClient(create_app(settings)) as client:
        body = execute_body(
            llm={"provider": "claude_code", "api_key": None},
            task={"title": "Sources", "kind": "research"},
        )
        events = parse_ndjson(client.post("/v1/execute", json=body, headers=AUTH).content)
    assert events[-1]["type"] == "result", events[-1]
    assert events[-1]["output"] == "# Sources\n\nA long plain answer."


def test_cancelling_a_turn_stops_the_cli(tmp_path: Path) -> None:
    """A node the backend gave up on must not leave `claude` generating in the background."""
    pid_file = tmp_path / "pid"
    cli = tmp_path / "claude"
    cli.write_text(
        "#!/usr/bin/env python3\n"
        "import os, sys, time\n"
        f"open({str(pid_file)!r}, 'w').write(str(os.getpid()))\n"
        "time.sleep(60)\n"
    )
    cli.chmod(cli.stat().st_mode | stat.S_IEXEC)

    async def scenario() -> int:
        provider = ClaudeCodeProvider(model="claude-sonnet-5", binary=str(cli))
        try:
            turn = asyncio.create_task(provider._invoke("system", "prompt", None, 1024))
            # The fake CLI reports its pid through a file; there is nothing to await on.
            while not pid_file.exists() or not pid_file.read_text():  # noqa: ASYNC110
                await asyncio.sleep(0.02)
            turn.cancel()
            with pytest.raises(asyncio.CancelledError):
                await turn
        finally:
            await provider.aclose()
        return int(pid_file.read_text())

    pid = asyncio.run(scenario())
    with pytest.raises(ProcessLookupError):
        os.kill(pid, 0)
