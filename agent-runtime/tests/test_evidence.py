"""Citation ids on what an agent is given, and the Sources section its deliverables keep."""

from __future__ import annotations

import io
from collections.abc import Callable
from pathlib import Path

import docx

from nexc_runtime.config import Settings
from nexc_runtime.evidence import Sources, find_citations, split_head
from nexc_runtime.streaming import EventStream
from nexc_runtime.tools import ToolContext, Workspace
from nexc_runtime.tools.agent_tools import FINISH
from nexc_runtime.tools.docx_tool import MAKE_DOCX
from nexc_runtime.tools.files import WRITE_FILE

CHUNK = "0199a1b2-0000-7000-8000-000000000001"
MEMORY = "0199a1b2-0000-7000-8000-000000000002"
NODE = "0199a1b2-0000-7000-8000-000000000003"


def sources() -> Sources:
    return Sources.from_context(
        upstream=[(NODE, "Findings"), ("", "no id, older backend")],
        memories=[f"[mem:{MEMORY}] The team writes in British English.", "no id either"],
        documents=[f"[doc:{CHUNK}] customs.pdf, p. 17\nGoods under 150 EUR are exempt."],
    )


def test_citations_are_read_outside_code_once_each() -> None:
    text = (
        f"Revenue grew [doc:{CHUNK}]. Costs fell. [mem:{MEMORY}] Both [doc:{CHUNK}, node:{NODE}].\n"
        f"See `[doc:{CHUNK}]` in code, [the report](http://x), [1] and [customs.pdf, p. 17].\n"
        f"```\n[node:{NODE}] not a citation\n```\nCertain [assumption]."
    )
    assert find_citations(text) == [f"doc:{CHUNK}", f"mem:{MEMORY}", f"node:{NODE}"]
    assert find_citations("nothing") == []
    assert split_head(f"[doc:{CHUNK}] customs.pdf, p. 17") == (f"doc:{CHUNK}", "customs.pdf, p. 17")
    assert split_head("[customs.pdf, p. 17] text") == ("", "[customs.pdf, p. 17] text")


def test_sources_are_read_off_the_context_and_name_what_is_cited() -> None:
    s = sources()
    assert set(s.items) == {f"node:{NODE}", f"mem:{MEMORY}", f"doc:{CHUNK}"}
    assert s.items[f"doc:{CHUNK}"].title == "customs.pdf, p. 17"
    assert s.items[f"mem:{MEMORY}"].title == "The team writes in British English."
    assert s.items[f"node:{NODE}"].title == "Findings"
    text = f"Exempt under 150 EUR [doc:{CHUNK}]. Built on the findings [node:{NODE}] [doc:nope]."
    cited = s.cited_in(text)
    assert [c.id for c in cited] == [f"doc:{CHUNK}", f"node:{NODE}", "doc:nope"]
    assert cited[2].title == "not among what the agent was given"
    section = s.markdown_section(text)
    assert section.startswith("## Sources\n\n- [doc:")
    assert f"- [node:{NODE}] Findings" in section
    with_sources = s.with_sources(text)
    assert with_sources.endswith(section)
    assert s.with_sources(with_sources) == with_sources, "appended once"
    assert s.with_sources("cites nothing") == "cites nothing"
    assert s.with_sources("# A\n\n## References\n\nx [doc:nope]").count("## ") == 1


async def test_deliverables_keep_a_sources_section(
    tmp_path: Path, make_settings: Callable[..., Settings]
) -> None:
    ctx = ToolContext(
        workspace=Workspace(tmp_path, max_file_bytes=100_000, max_total_bytes=1_000_000),
        events=EventStream(),
        settings=make_settings(),
        agent_name="writer",
        sources=sources(),
    )
    # finish: the final answer names its sources.
    await FINISH.invoke(ctx, {"answer": f"Goods under 150 EUR are exempt [doc:{CHUNK}]."})
    assert ctx.final_answer is not None
    assert f"## Sources\n\n- [doc:{CHUNK}] customs.pdf, p. 17" in ctx.final_answer
    # write_file: a Markdown file does too; other files are left alone.
    await WRITE_FILE.invoke(ctx, {"path": "notes.md", "content": f"Exempt [doc:{CHUNK}]."})
    assert "## Sources" in (tmp_path / "notes.md").read_text()
    await WRITE_FILE.invoke(ctx, {"path": "notes.txt", "content": f"Exempt [doc:{CHUNK}]."})
    assert "## Sources" not in (tmp_path / "notes.txt").read_text()
    # make_docx: a Sources section lists the cited ids with their titles.
    await MAKE_DOCX.invoke(
        ctx,
        {
            "filename": "report.docx",
            "title": "Customs",
            "sections": [
                {
                    "heading": "Findings",
                    "paragraphs": [f"Goods under 150 EUR are exempt [doc:{CHUNK}]."],
                    "bullets": [f"British English throughout [mem:{MEMORY}]"],
                    "table_rows": [],
                }
            ],
        },
    )
    document = docx.Document(io.BytesIO((tmp_path / "report.docx").read_bytes()))
    text = [p.text for p in document.paragraphs]
    assert "Sources" in text
    assert f"[doc:{CHUNK}] customs.pdf, p. 17" in text
    assert f"[mem:{MEMORY}] The team writes in British English." in text
    # A document that cites nothing gets no section.
    await MAKE_DOCX.invoke(
        ctx,
        {
            "filename": "plain.docx",
            "title": "Plain",
            "sections": [{"heading": "A", "paragraphs": ["x"], "bullets": [], "table_rows": []}],
        },
    )
    plain = docx.Document(io.BytesIO((tmp_path / "plain.docx").read_bytes()))
    assert "Sources" not in [p.text for p in plain.paragraphs]
