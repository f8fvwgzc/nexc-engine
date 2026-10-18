"""The offline demo provider drives the real tool loop end to end."""

from __future__ import annotations

import base64
import io
from collections.abc import Callable
from typing import Any

import docx
import pytest
from fastapi.testclient import TestClient

from nexc_runtime.api.app import create_app
from nexc_runtime.config import Settings
from nexc_runtime.llm.base import ToolResult, ToolSpec
from nexc_runtime.llm.demo import DEMO_BANNER, DemoBrief, DemoProvider

from .conftest import AUTH, execute_body, parse_ndjson


def _run(make_settings: Callable[..., Settings], **overrides: Any) -> list[dict[str, Any]]:
    with TestClient(create_app(make_settings())) as client:
        body = execute_body(llm={"provider": "demo", "api_key": None}, **overrides)
        response = client.post("/v1/execute", json=body, headers=AUTH)
    assert response.status_code == 200
    return parse_ndjson(response.content)


def test_demo_document_produces_real_docx(make_settings: Callable[..., Settings]) -> None:
    events = _run(make_settings)
    assert events[-1]["type"] == "result"
    assert events[-1]["output"].startswith(DEMO_BANNER)
    assert events[-1]["tokens_in"] > 0 and events[-1]["tokens_out"] > 0
    assert any(e["type"] == "delta" for e in events)
    (artifact,) = [e for e in events if e["type"] == "artifact"]
    assert artifact["path"] == "quarterly-report.docx"
    document = docx.Document(io.BytesIO(base64.b64decode(artifact["content_b64"])))
    assert document.paragraphs[0].text == "Quarterly report (demo)"
    assert len(document.tables) == 1


def test_demo_output_compiles_upstream_into_docx(make_settings: Callable[..., Settings]) -> None:
    upstream = [
        {
            "node_id": "a",
            "title": "Findings",
            "output": f"{DEMO_BANNER}\n\n## Findings\nRevenue grew 12%.",
        },
        {"node_id": "b", "title": "Risks", "output": "- Supply chain delays"},
    ]
    events = _run(
        make_settings,
        task={"title": "Final report", "kind": "output"},
        context={"upstream": upstream},
    )
    (artifact,) = [e for e in events if e["type"] == "artifact"]
    assert artifact["path"] == "final-report.docx"
    document = docx.Document(io.BytesIO(base64.b64decode(artifact["content_b64"])))
    text = [p.text for p in document.paragraphs]
    assert "Findings" in text and "Revenue grew 12%." in text
    assert "Risks" in text and "Supply chain delays" in text
    assert not any("demo mode - generated offline" in t.lower() for t in text[1:])


@pytest.mark.parametrize(
    ("kind", "expected"),
    [("code", ["quarterly_report.py"]), ("research", ["quarterly-report-notes.md"]), ("topic", [])],
)
def test_demo_other_kinds(
    make_settings: Callable[..., Settings], kind: str, expected: list[str]
) -> None:
    events = _run(make_settings, task={"kind": kind})
    assert events[-1]["type"] == "result"
    assert [e["path"] for e in events if e["type"] == "artifact"] == expected


FETCH = ToolSpec("fetch_url", "Fetch a URL through the gateway.", {"type": "object"})
WRITE = ToolSpec("write_file", "Write a file.", {"type": "object"})
FIGURES = "http://127.0.0.1:8099/quarterly-figures.json"


async def test_demo_reads_the_briefs_urls_through_the_gateway_when_offered() -> None:
    brief = DemoBrief(
        title="Quarterly numbers",
        kind="research",
        content=f"Read the figures from {FIGURES}, then {FIGURES} again (see {FIGURES}).",
        role="analyst",
    )
    provider = DemoProvider(brief)
    conversation = provider.new_conversation("system", "go")
    first = await provider.next_turn(conversation, [FETCH, WRITE], max_tokens=100)
    assert first.stop == "tool_use"
    assert [(c.name, c.input) for c in first.tool_calls] == [("fetch_url", {"url": FIGURES})]
    assert FIGURES in first.text
    # The gateway refused it; the demo goes on to its deliverable all the same.
    provider.add_tool_results(
        conversation, [ToolResult(first.tool_calls[0].id, "not allowed", is_error=True)]
    )
    second = await provider.next_turn(conversation, [FETCH, WRITE], max_tokens=100)
    assert [c.name for c in second.tool_calls] == ["write_file"]
    assert FIGURES in second.tool_calls[0].input["content"]
    provider.add_tool_results(conversation, [ToolResult(second.tool_calls[0].id, "written")])
    third = await provider.next_turn(conversation, [FETCH, WRITE], max_tokens=100)
    assert third.stop == "end_turn" and third.tool_calls == []
    # Without the gateway tool (the NDJSON wire), nothing is fetched and the run is as before.
    offline = DemoProvider(brief)
    conversation = offline.new_conversation("system", "go")
    turn = await offline.next_turn(conversation, [WRITE], max_tokens=100)
    assert [c.name for c in turn.tool_calls] == ["write_file"]


CHUNK = "0199a1b2-0000-7000-8000-000000000011"
MEMORY = "0199a1b2-0000-7000-8000-000000000012"
NODE = "0199a1b2-0000-7000-8000-000000000013"


def test_demo_cites_what_it_was_given_and_keeps_a_sources_section(
    make_settings: Callable[..., Settings],
) -> None:
    context = {
        "upstream": [{"node_id": NODE, "title": "Findings", "output": "Revenue grew 12%."}],
        "memories": [f"[mem:{MEMORY}] The team writes in British English. Always."],
        "documents": [
            f"[doc:{CHUNK}] customs.pdf, p. 17\nGoods under 150 EUR are exempt from duty.\nMore."
        ],
    }
    events = _run(make_settings, context=context)
    output = events[-1]["output"]
    assert f"Goods under 150 EUR are exempt from duty. [doc:{CHUNK}] (customs.pdf, p. 17)" in output
    assert f"Recalled: The team writes in British English. [mem:{MEMORY}]" in output
    assert f"Built on “Findings”: its findings feed into this step [node:{NODE}]" in output
    assert "[assumption]" in output
    sources = output.split("## Sources", 1)[1]
    assert f"- [doc:{CHUNK}] customs.pdf, p. 17" in sources
    assert f"- [node:{NODE}] Findings" in sources
    (artifact,) = [e for e in events if e["type"] == "artifact"]
    document = docx.Document(io.BytesIO(base64.b64decode(artifact["content_b64"])))
    text = [p.text for p in document.paragraphs]
    assert "Evidence" in text and "Sources" in text
    assert any(f"[doc:{CHUNK}]" in t for t in text)
    # Given nothing, the answer declares every claim its own and cites nothing.
    bare = _run(make_settings)[-1]["output"]
    assert "[doc:" not in bare and "## Sources" not in bare
    assert bare.count("[assumption]") == 3
    # The final deliverable cites each upstream node its sections rest on.
    compiled = _run(
        make_settings,
        task={"title": "Final report", "kind": "output"},
        context={"upstream": [{"node_id": NODE, "title": "Findings", "output": "Revenue grew."}]},
    )
    (artifact,) = [e for e in compiled if e["type"] == "artifact"]
    document = docx.Document(io.BytesIO(base64.b64decode(artifact["content_b64"])))
    text = [p.text for p in document.paragraphs]
    assert "Revenue grew." in text
    assert f"Rests on the upstream result “Findings” [node:{NODE}]." in text
    assert f"[node:{NODE}] Findings" in text


def test_demo_names_at_most_three_urls_once_each() -> None:
    urls = [f"https://example.com/{n}." for n in range(5)]
    brief = DemoBrief(title="t", kind="research", content=" ".join(urls + urls))
    assert DemoProvider(brief).urls() == [u.rstrip(".") for u in urls[:3]]
    assert DemoProvider(DemoBrief(title="t", kind="research", content="no links")).urls() == []


def _comparable(events: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Events with .docx bytes replaced by their text (the files embed creation timestamps) and
    span ids, clocks and durations left out (they are what make each run its own)."""
    out = []
    for event in events:
        if event["type"] == "span":
            attributes = {k: v for k, v in event["attributes"].items() if k != "duration_ns"}
            clocks = ("id", "parent_id", "started_at", "ended_at")
            kept = {k: v for k, v in event.items() if k not in clocks}
            linked = event["parent_id"] is not None
            out.append({**kept, "attributes": attributes, "has_parent": linked})
            continue
        if event["type"] == "artifact" and event["path"].endswith(".docx"):
            document = docx.Document(io.BytesIO(base64.b64decode(event["content_b64"])))
            text = [p.text for p in document.paragraphs]
            out.append({**event, "content_b64": text, "size": None})
        else:
            out.append(event)
    return out


def test_demo_is_deterministic(make_settings: Callable[..., Settings]) -> None:
    first = _run(make_settings)
    second = _run(make_settings)
    assert _comparable(first) == _comparable(second)
