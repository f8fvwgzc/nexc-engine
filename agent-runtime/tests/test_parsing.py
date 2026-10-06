"""Document parsing: every format's block order, tables and headers, and the /v1/parse route."""

from __future__ import annotations

import io
import threading
from collections.abc import Callable
from typing import Any
from urllib.parse import quote

import docx
import openpyxl
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.platypus import Table as PdfTable
from reportlab.platypus import TableStyle
from starlette.middleware.base import BaseHTTPMiddleware

from nexc_runtime.parsing import (
    Block,
    UnparsableDocumentError,
    UnsupportedFormatError,
    detect_format,
    parse_document,
)
from nexc_runtime.parsing import xlsx as xlsx_parser
from nexc_runtime.parsing.tables import estimate_header_rows, table_block

from .conftest import AUTH, FakeProvider

ClientFactory = Callable[..., tuple[TestClient, FakeProvider]]

REGION_TABLE = [
    ["Region", "2024", "2024"],
    ["Region", "Q1", "Q2"],
    ["EMEA", "10", "12"],
    ["APAC", "7", "9"],
]


def kinds(blocks: list[Block]) -> list[str]:
    return [block.kind for block in blocks]


def assert_rectangular(block: Block) -> None:
    assert block.rows
    assert len({len(row) for row in block.rows}) == 1
    assert block.text == ""


# --- fixtures built in memory ---------------------------------------------------------------


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("Revenue   grew\tin every region.")
    table = document.add_table(rows=4, cols=3)
    for row, values in zip(table.rows, REGION_TABLE, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    table.cell(0, 1).merge(table.cell(0, 2)).text = "2024"  # header cell spanning two columns
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Region"  # and one spanning two rows
    document.add_heading("Outlook", level=2)
    document.add_paragraph("")
    document.add_paragraph("Stable.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(extra_rows: int = 0) -> bytes:
    workbook = openpyxl.Workbook()
    sales = workbook.active
    assert sales is not None
    sales.title = "Sales"
    sales.append(["Region", "2024", None])
    sales.append([None, "Q1", "Q2"])
    sales.append(["EMEA", 10, 12.5])
    sales.append([None, None, None])
    sales.append(["APAC", 7.0, True])
    for index in range(extra_rows):
        sales.append([f"R{index}", index, index])
    workbook.create_sheet("Empty")
    notes = workbook.create_sheet("Notes")
    notes["B2"] = "Owner"
    notes["C2"] = "Status"
    notes["B3"] = "Ann"
    notes["C3"] = "done"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    assert first.shapes.title is not None
    first.shapes.title.text = "Roadmap"
    body = first.placeholders[1].text_frame
    body.text = "Ship the parser"
    body.add_paragraph().text = "Index the blocks"

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    assert second.shapes.title is not None
    second.shapes.title.text = "Numbers"
    frame = second.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(6), Inches(2))
    table = frame.table
    for row_index, values in enumerate([REGION_TABLE[0], REGION_TABLE[1], REGION_TABLE[2]]):
        for column, value in enumerate(values):
            table.cell(row_index, column).text = value
    table.cell(0, 2).text = ""
    table.cell(0, 1).merge(table.cell(0, 2))
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_pdf(*, password: str | None = None, pages: int = 1) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, encrypt=password)
    width, height = A4
    for _ in range(pages):
        pdf.setFont("Helvetica-Bold", 22)
        pdf.drawString(72, height - 80, "Quarterly report")
        pdf.setFont("Helvetica", 10)
        pdf.drawString(72, height - 120, "Revenue grew in every region during the quar-")
        pdf.drawString(72, height - 132, "ter, and margins held.")
        pdf.drawString(72, height - 170, "A second paragraph follows after a gap.")
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(72, height - 210, "Regional breakdown")
        table = PdfTable([REGION_TABLE[0], ["", "Q1", "Q2"], *REGION_TABLE[2:]], colWidths=90)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                    ("SPAN", (1, 0), (2, 0)),
                    ("SPAN", (0, 0), (0, 1)),
                    ("FONTSIZE", (0, 0), (-1, -1), 10),
                ]
            )
        )
        table.wrapOn(pdf, width, height)
        table.drawOn(pdf, 72, height - 320)
        pdf.setFont("Helvetica", 10)
        pdf.drawString(72, height - 360, "Text after the table.")
        pdf.showPage()
    pdf.showPage()  # a page without a text layer, like a scan
    pdf.save()
    return buffer.getvalue()


HTML = b"""<!doctype html>
<html><head><title>ignored</title><style>p { color: red }</style></head>
<body>
<nav><a href="/">Home</a><p>Menu</p></nav>
<h1>Quarterly   report</h1>
<p>Revenue <b>grew</b><br>in every region &amp; segment.</p>
<script>var ignored = "<p>nope</p>";</script>
<ul><li>First point</li><li>Second point</li></ul>
<table>
  <caption>Sales by region</caption>
  <thead>
    <tr><th rowspan="2">Region</th><th colspan="2">2024</th></tr>
    <tr><th>Q1</th><th>Q2</th></tr>
  </thead>
  <tbody>
    <tr><th>EMEA</th><td>10</td><td>12</td></tr>
    <tr><td>APAC</td><td colspan="2">n/a</td></tr>
    <tr><td>LATAM</td></tr>
  </tbody>
</table>
<h3>Outlook</h3>
Loose text outside a paragraph.
</body></html>
"""

MARKDOWN = b"""---
title: ignored front matter
---
# Quarterly report

Revenue grew
in every region.

Setext title
------------

- First point
- Second point
  continued

| Region | Q1 | Q2 |
|:-------|---:|---:|
| EMEA   | 10 | 12 |
| A\\|B   | 7  |

```python
print("code")
```
###### Deep ######
"""


# --- formats --------------------------------------------------------------------------------


def test_docx_blocks_keep_document_order_and_repeat_merged_cells() -> None:
    parsed = parse_document("report.docx", make_docx())
    assert parsed.pages is None
    assert kinds(parsed.blocks) == ["heading", "text", "table", "heading", "text"]
    title, intro, table, outlook, closing = parsed.blocks
    assert (title.text, title.level) == ("Quarterly report", 1)
    assert intro.text == "Revenue grew in every region."  # whitespace collapsed
    assert (outlook.text, outlook.level) == ("Outlook", 2)
    assert closing.text == "Stable."  # the empty paragraph was dropped
    assert table.rows == REGION_TABLE  # "2024" and "Region" repeated across their spans
    assert table.header_rows == 2
    assert table.level is None and table.page is None
    assert_rectangular(table)


def test_docx_title_style_is_a_heading() -> None:
    document = docx.Document()
    document.add_heading("The title", level=0)  # python-docx: level 0 is the "Title" style
    buffer = io.BytesIO()
    document.save(buffer)
    (block,) = parse_document("t.docx", buffer.getvalue()).blocks
    assert (block.kind, block.level, block.text) == ("heading", 1, "The title")


def test_xlsx_one_table_per_non_empty_sheet() -> None:
    parsed = parse_document("book.xlsx", make_xlsx())
    assert parsed.pages == 3  # the empty sheet still counts
    assert kinds(parsed.blocks) == ["heading", "table", "heading", "table"]
    sales_title, sales, notes_title, notes = parsed.blocks
    assert (sales_title.text, sales_title.level, sales_title.page) == ("Sales", 2, 1)
    assert sales.rows == [
        ["Region", "2024", ""],  # merged cells are not exposed in read-only mode
        ["", "Q1", "Q2"],
        ["EMEA", "10", "12.5"],
        ["APAC", "7", "TRUE"],  # the blank row is gone, 7.0 reads as 7
    ]
    assert sales.header_rows == 2
    assert sales.page == 1
    assert (notes_title.text, notes_title.page) == ("Notes", 3)
    assert notes.rows == [["Owner", "Status"], ["Ann", "done"]]  # empty row / column trimmed
    assert notes.header_rows == 1


def test_xlsx_rows_are_capped_and_truncation_is_reported(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx_parser, "MAX_ROWS_PER_SHEET", 6)
    parsed = parse_document("book.xlsx", make_xlsx(extra_rows=10))
    sales = parsed.blocks[1]
    assert sales.rows is not None and len(sales.rows) == 6
    note = parsed.blocks[2]
    assert note.kind == "text" and "truncated" in note.text and note.page == 1


def test_pptx_title_then_text_and_tables_per_slide() -> None:
    parsed = parse_document("deck.pptx", make_pptx())
    assert parsed.pages == 2
    assert kinds(parsed.blocks) == ["heading", "text", "text", "heading", "table"]
    assert [(b.text, b.page) for b in parsed.blocks[:4]] == [
        ("Roadmap", 1),
        ("Ship the parser", 1),
        ("Index the blocks", 1),
        ("Numbers", 2),
    ]
    assert parsed.blocks[0].level == 2
    table = parsed.blocks[4]
    assert table.rows == REGION_TABLE[:3]  # merged "2024" repeated
    assert table.header_rows == 2
    assert table.page == 2


def test_html_blocks_tables_and_ignored_elements() -> None:
    parsed = parse_document("page.html", HTML)
    assert parsed.pages is None
    expected = ["heading", "text", "text", "text", "text", "table", "heading", "text"]
    assert kinds(parsed.blocks) == expected
    texts = [block.text for block in parsed.blocks]
    assert texts[0] == "Quarterly report" and parsed.blocks[0].level == 1
    assert texts[1] == "Revenue grew in every region & segment."
    assert texts[2:5] == ["First point", "Second point", "Sales by region"]
    assert (texts[6], parsed.blocks[6].level) == ("Outlook", 3)
    assert texts[7] == "Loose text outside a paragraph."
    assert not any("Menu" in t or "nope" in t or "ignored" in t or "color" in t for t in texts)
    table = parsed.blocks[5]
    assert table.rows == [
        ["Region", "2024", "2024"],  # colspan repeated
        ["Region", "Q1", "Q2"],  # rowspan repeated
        ["EMEA", "10", "12"],
        ["APAC", "n/a", "n/a"],
        ["LATAM", "", ""],  # short row padded
    ]
    assert table.header_rows == 2  # the <th> rows; the row-header <th> of EMEA does not count
    assert_rectangular(table)


def test_html_table_without_th_falls_back_to_the_heuristic() -> None:
    html = b"<table><tr><td>Name</td><td>Qty</td></tr><tr><td>Bolt</td><td>4</td></tr></table>"
    (table,) = parse_document("t.htm", html).blocks
    assert table.rows == [["Name", "Qty"], ["Bolt", "4"]]
    assert table.header_rows == 1


def test_csv_and_tsv_are_one_table() -> None:
    (table,) = parse_document("data.csv", b'name;note\r\nBolt;"a; b"\r\n\r\nNut\r\n').blocks
    assert table.rows == [["name", "note"], ["Bolt", "a; b"], ["Nut", ""]]
    assert table.header_rows == 1
    (table,) = parse_document("data.tsv", b"a\tb\n1\t2, 3\n").blocks
    assert table.rows == [["a", "b"], ["1", "2, 3"]]
    assert parse_document("empty.csv", b"\n\n").blocks == []


def test_markdown_headings_paragraphs_lists_and_pipe_tables() -> None:
    parsed = parse_document("notes.md", MARKDOWN)
    assert [(b.kind, b.level, b.text) for b in parsed.blocks if b.kind != "table"] == [
        ("heading", 1, "Quarterly report"),
        ("text", None, "Revenue grew in every region."),
        ("heading", 2, "Setext title"),
        ("text", None, "First point"),
        ("text", None, "Second point continued"),
        ("text", None, 'print("code")'),
        ("heading", 6, "Deep"),
    ]
    assert kinds(parsed.blocks)[5] == "table"
    table = parsed.blocks[5]
    assert table.rows == [["Region", "Q1", "Q2"], ["EMEA", "10", "12"], ["A|B", "7", ""]]
    assert table.header_rows == 1


def test_plain_text_splits_on_blank_lines_with_latin1_fallback() -> None:
    data = "First  line\nsame paragraph.\r\n\r\n \r\nCaf\u00e9 second.\n".encode("latin-1")
    parsed = parse_document("notes.txt", data)
    assert [b.text for b in parsed.blocks] == ["First line same paragraph.", "Caf\u00e9 second."]
    assert all(b.kind == "text" and b.page is None for b in parsed.blocks)


def test_pdf_headings_paragraphs_and_tables_in_reading_order() -> None:
    parsed = parse_document("report.pdf", make_pdf())
    assert parsed.pages == 2  # the page without text is counted but yields nothing
    assert [(b.kind, b.level, b.text) for b in parsed.blocks] == [
        ("heading", 1, "Quarterly report"),
        ("text", None, "Revenue grew in every region during the quarter, and margins held."),
        ("text", None, "A second paragraph follows after a gap."),
        ("heading", 4, "Regional breakdown"),
        ("table", None, ""),
        ("text", None, "Text after the table."),
    ]
    assert all(block.page == 1 for block in parsed.blocks)
    table = parsed.blocks[4]
    assert table.rows == REGION_TABLE  # merged cells repeated, cell text not emitted twice
    assert table.header_rows == 2


# --- detection and failures -----------------------------------------------------------------


def test_format_is_sniffed_when_the_extension_does_not_say() -> None:
    assert detect_format("upload", make_pdf()) == "pdf"
    assert detect_format("upload.bin", make_docx()) == "docx"
    assert detect_format("", make_xlsx()) == "xlsx"
    assert detect_format("deck", make_pptx()) == "pptx"
    assert detect_format("config.yaml", b"key: value\n") == "text"
    assert detect_format("main.rs", b"fn main() {}\n") == "text"
    assert detect_format("C:\\docs\\Report.PDF", b"") == "pdf"


@pytest.mark.parametrize(
    ("filename", "data"),
    [
        ("photo.png", b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR" + bytes(64)),
        ("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(64)),
        ("archive.zip", b"PK\x03\x04" + bytes(64)),
    ],
)
def test_unsupported_types_are_rejected(filename: str, data: bytes) -> None:
    with pytest.raises(UnsupportedFormatError):
        parse_document(filename, data)


@pytest.mark.parametrize(
    ("filename", "data"),
    [
        ("broken.pdf", b"%PDF-1.7\nthis is not a pdf"),
        ("broken.docx", b"PK\x03\x04 not a zip"),
        ("broken.xlsx", b"not a workbook"),
        ("broken.pptx", b"not a deck"),
        ("empty.txt", b""),
    ],
)
def test_corrupt_files_are_unparsable(filename: str, data: bytes) -> None:
    with pytest.raises(UnparsableDocumentError):
        parse_document(filename, data)


def test_encrypted_pdf_is_unparsable_with_a_clear_message() -> None:
    with pytest.raises(UnparsableDocumentError, match="password"):
        parse_document("secret.pdf", make_pdf(password="hunter2"))


# --- table heuristics -----------------------------------------------------------------------


@pytest.mark.parametrize(
    ("rows", "expected"),
    [
        ([], 0),
        ([["Name", "Qty"]], 1),
        ([["Name", "Qty"], ["Bolt", "4"]], 1),
        ([["Name", "Role"], ["Ann", "Dev"], ["Bob", ""]], 1),  # text-only data stays data
        (REGION_TABLE, 2),  # merged label repeated
        ([["Region", "2024", ""], ["", "Q1", "Q2"], ["EMEA", "1,5", "2"]], 2),  # not repeated
        ([["Region", "2024", ""], ["", "Q1", "Q2"], ["EMEA", "high", "low"]], 2),
        (
            [
                ["", "Revenue", "Revenue", "Revenue", "Costs"],
                ["", "2023", "2024", "2024", "2024"],
                ["Unit", "FY", "H1", "H2", "FY"],
                ["North", "1 200", "(35)", "12%", "$4.5"],
                ["South", "900", "40", "3%", "$2"],
            ],
            3,
        ),
        ([["A", "A"], ["x", "y"], ["z", "w"], ["1", "2"]], 2),
    ],
)
def test_header_rows_are_estimated(rows: list[list[str]], expected: int) -> None:
    assert estimate_header_rows(rows) == expected


def test_table_block_pads_cleans_and_drops_empty_tables() -> None:
    block = table_block([["a  b", None], ["c"], []])
    assert block is not None
    assert block.rows == [["a b", ""], ["c", ""], ["", ""]]
    assert block.header_rows == 1
    assert table_block([["", None], [" "]]) is None
    compact = table_block([["", "a", ""], ["", "", ""], ["", "b", ""]], compact=True)
    assert compact is not None and compact.rows == [["a"], ["b"]]


# --- HTTP route -----------------------------------------------------------------------------


def post(client: TestClient, filename: str, data: bytes, **kwargs: Any) -> Any:
    headers = {"Content-Type": "application/octet-stream", **kwargs.pop("headers", AUTH)}
    return client.post(f"/v1/parse?filename={quote(filename)}", content=data, headers=headers)


def test_parse_route_returns_the_contract_shape(client_for: ClientFactory) -> None:
    client, _ = client_for([])
    response = post(client, "Q3 report \u2013 final.docx", make_docx())
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/json"
    body = response.json()
    assert set(body) == {"pages", "blocks"}
    assert body["pages"] is None
    for block in body["blocks"]:
        assert set(block) == {"kind", "level", "text", "page", "rows", "header_rows"}
    assert body["blocks"][0] == {
        "kind": "heading",
        "level": 1,
        "text": "Quarterly report",
        "page": None,
        "rows": None,
        "header_rows": None,
    }
    assert body["blocks"][2] == {
        "kind": "table",
        "level": None,
        "text": "",
        "page": None,
        "rows": REGION_TABLE,
        "header_rows": 2,
    }


def test_parse_route_reports_pages(client_for: ClientFactory) -> None:
    client, _ = client_for([])
    body = post(client, "report.pdf", make_pdf()).json()
    assert body["pages"] == 2
    assert body["blocks"][0]["page"] == 1


def test_parse_route_sniffs_without_a_filename(client_for: ClientFactory) -> None:
    client, _ = client_for([])
    response = client.post("/v1/parse", content=make_xlsx(), headers=AUTH)
    assert response.status_code == 200
    assert response.json()["pages"] == 3


@pytest.mark.parametrize(
    "headers",
    [
        {},
        {"Authorization": "Bearer wrong-token-wrong-token-wrong-token!!"},
        {"Authorization": "Basic " + "t" * 40},
    ],
)
def test_parse_route_requires_valid_bearer(
    client_for: ClientFactory, headers: dict[str, str]
) -> None:
    client, _ = client_for([])
    response = post(client, "notes.txt", b"hello", headers=headers)
    assert response.status_code == 401
    assert response.headers["www-authenticate"] == "Bearer"


def test_parse_route_415_for_unsupported_types(client_for: ClientFactory) -> None:
    client, _ = client_for([])
    response = post(client, "photo.png", b"\x89PNG\r\n\x1a\n\x00\x00" + bytes(32))
    assert response.status_code == 415
    assert response.json() == {"detail": '".png" files cannot be parsed'}


def test_parse_route_422_for_unparsable_files(client_for: ClientFactory) -> None:
    client, _ = client_for([])
    response = post(client, "broken.pdf", b"%PDF-1.7\nSECRET-CONTENT garbage")
    assert response.status_code == 422
    assert response.json() == {"detail": "the PDF is corrupt or not a PDF"}
    response = post(client, "secret.pdf", make_pdf(password="hunter2"))
    assert response.status_code == 422
    assert "password" in response.json()["detail"]
    assert post(client, "empty.docx", b"").status_code == 422


def test_parse_route_413_above_its_own_limit(client_for: ClientFactory) -> None:
    client, _ = client_for([], runtime_max_parse_bytes=4096, runtime_max_request_bytes=1024)
    # larger than the JSON limit of the other routes, within the parse limit
    assert post(client, "notes.txt", b"x" * 3000).status_code == 200
    response = post(client, "notes.txt", b"x" * 5000)
    assert response.status_code == 413
    assert "4096" in response.json()["detail"]


def test_parse_limit_defaults_to_50_mib(settings: Any) -> None:
    assert settings.runtime_max_parse_bytes == 50 * 1024 * 1024


def test_parsing_runs_off_the_event_loop(
    client_for: ClientFactory, monkeypatch: pytest.MonkeyPatch
) -> None:
    parser_threads: list[int] = []
    loop_threads: list[int] = []
    real = parse_document

    def spy(filename: str, data: bytes) -> Any:
        parser_threads.append(threading.get_ident())
        return real(filename, data)

    monkeypatch.setattr("nexc_runtime.api.app.parse_document", spy)
    client, _ = client_for([])

    async def record(request: Any, call_next: Any) -> Any:
        loop_threads.append(threading.get_ident())
        return await call_next(request)

    client.app.add_middleware(BaseHTTPMiddleware, dispatch=record)  # type: ignore[attr-defined]
    assert post(client, "notes.txt", b"hello").status_code == 200
    assert len(parser_threads) == 1 and len(loop_threads) == 1
    assert parser_threads != loop_threads
