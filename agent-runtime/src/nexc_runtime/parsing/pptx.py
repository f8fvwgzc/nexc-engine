"""PowerPoint (.pptx) via python-pptx: per slide the title, then text frames and tables."""

from __future__ import annotations

import io
from collections.abc import Iterable, Iterator

from pptx import Presentation
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.table import Table

from .model import Block, ParsedDocument, heading, paragraph
from .tables import table_block

SLIDE_HEADING_LEVEL = 2


def _walk(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _walk(shape.shapes)
        else:
            yield shape


def _table_rows(table: Table) -> list[list[str]]:
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    for row_index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            if not cell.is_merge_origin:
                continue
            for down in range(cell.span_height):
                for across in range(cell.span_width):
                    if row_index + down < len(rows) and column + across < len(rows[0]):
                        rows[row_index + down][column + across] = cell.text
    return rows


def parse_pptx(data: bytes) -> ParsedDocument:
    presentation = Presentation(io.BytesIO(data))
    blocks: list[Block | None] = []
    slides = list(presentation.slides)
    for number, slide in enumerate(slides, start=1):
        title = slide.shapes.title
        if title is not None:
            blocks.append(heading(title.text_frame.text, SLIDE_HEADING_LEVEL, number))
        for shape in _walk(slide.shapes):
            if title is not None and shape.shape_id == title.shape_id:
                continue
            if isinstance(shape, GraphicFrame) and shape.has_table:
                blocks.append(table_block(_table_rows(shape.table), page=number))
            elif shape.has_text_frame:
                frame = shape.text_frame  # type: ignore[attr-defined]
                blocks += [paragraph(item.text, number) for item in frame.paragraphs]
    return ParsedDocument(
        pages=len(slides), blocks=[block for block in blocks if block is not None]
    )
