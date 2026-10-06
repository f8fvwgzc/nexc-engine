"""Word (.docx) via python-docx: headings by style, paragraphs and tables in document order."""

from __future__ import annotations

import io
import re

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from .model import Block, ParsedDocument, heading, paragraph
from .tables import table_block

_HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)
_NAMED_LEVELS = {"title": 1, "subtitle": 2}


def _heading_level(item: Paragraph) -> int | None:
    style = item.style
    name = (style.name if style is not None else None) or ""
    if match := _HEADING_STYLE.match(name.strip()):
        return int(match.group(1))
    return _NAMED_LEVELS.get(name.strip().lower())


def _table_rows(table: Table) -> list[list[str]]:
    # `row.cells` yields the same cell for every grid position of a merged span (both
    # directions), so merged values come out repeated.
    return [[cell.text for cell in row.cells] for row in table.rows]


def parse_docx(data: bytes) -> ParsedDocument:
    document = docx.Document(io.BytesIO(data))
    blocks: list[Block | None] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            blocks.append(table_block(_table_rows(item)))
            continue
        level = _heading_level(item)
        blocks.append(heading(item.text, level) if level else paragraph(item.text))
    return ParsedDocument(blocks=[block for block in blocks if block is not None])
