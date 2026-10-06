"""Excel (.xlsx / .xlsm) via openpyxl in read-only mode: one table per non-empty sheet."""

from __future__ import annotations

import datetime as dt
import io
from typing import Any

import openpyxl

from .model import Block, ParsedDocument, heading, paragraph
from .tables import table_block

MAX_ROWS_PER_SHEET = 5000
SHEET_HEADING_LEVEL = 2


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, dt.datetime):
        at_midnight = value.time() == dt.time(0)
        return value.date().isoformat() if at_midnight else value.isoformat(sep=" ")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    return str(value)


def _sheet_rows(sheet: Any) -> tuple[list[list[str]], bool]:
    """Non-empty rows of a sheet (trailing blanks trimmed) and whether the cap cut it short."""
    # The stored dimensions are optional and often wrong; without them openpyxl just streams.
    if hasattr(sheet, "reset_dimensions"):
        sheet.reset_dimensions()
    rows: list[list[str]] = []
    for values in sheet.iter_rows(values_only=True):
        row = [_cell_text(value) for value in values]
        while row and not row[-1].strip():
            row.pop()
        if not row:
            continue
        if len(rows) == MAX_ROWS_PER_SHEET:
            return rows, True
        rows.append(row)
    return rows, False


def parse_xlsx(data: bytes) -> ParsedDocument:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets: list[Any] = workbook.worksheets  # chart sheets are not included
        blocks: list[Block | None] = []
        for number, sheet in enumerate(sheets, start=1):
            rows, truncated = _sheet_rows(sheet)
            table = table_block(rows, page=number, compact=True)
            if table is None:
                continue
            blocks += [heading(str(sheet.title), SHEET_HEADING_LEVEL, number), table]
            if truncated:
                blocks.append(
                    paragraph(
                        f'Sheet "{sheet.title}" was truncated to its first '
                        f"{MAX_ROWS_PER_SHEET:,} rows.",
                        number,
                    )
                )
        return ParsedDocument(
            pages=len(sheets), blocks=[block for block in blocks if block is not None]
        )
    finally:
        workbook.close()
