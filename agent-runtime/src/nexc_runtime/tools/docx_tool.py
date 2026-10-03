"""make_docx: turn structured content into a real Word document with python-docx."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pydantic import Field

from .base import Tool, ToolArgs, ToolContext, ToolError

MAX_SECTIONS = 200
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 20


class DocxSection(ToolArgs):
    heading: str = Field(description="Section heading.")
    paragraphs: list[str] = Field(description="Body paragraphs, in order. May be empty.")
    bullets: list[str] = Field(
        description="Bullet points shown after the paragraphs. May be empty."
    )
    table_rows: list[list[str]] = Field(
        description="Optional table after the bullets: first row is the header. Empty for none."
    )


class MakeDocxArgs(ToolArgs):
    filename: str = Field(description="Relative output path ending in .docx, e.g. 'report.docx'.")
    title: str = Field(description="Document title.")
    sections: list[DocxSection] = Field(description="Sections in reading order.")


def render_docx(args: MakeDocxArgs) -> bytes:
    """Render the document to bytes (pure function, easy to test)."""
    if len(args.sections) > MAX_SECTIONS:
        raise ToolError(f"too many sections (max {MAX_SECTIONS})")
    document = Document()
    document.core_properties.title = args.title
    document.core_properties.author = "nexc-engine"
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    document.add_heading(args.title, level=0)
    for section in args.sections:
        document.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            body = document.add_paragraph(paragraph)
            body.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for bullet in section.bullets:
            document.add_paragraph(bullet, style="List Bullet")
        if section.table_rows:
            _add_table(document, section.table_rows)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_table(document: DocxDocument, rows: list[list[str]]) -> None:
    if len(rows) > MAX_TABLE_ROWS:
        raise ToolError(f"table has too many rows (max {MAX_TABLE_ROWS})")
    width = max((len(r) for r in rows), default=0)
    if width == 0:
        return
    if width > MAX_TABLE_COLS:
        raise ToolError(f"table has too many columns (max {MAX_TABLE_COLS})")
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            cell.text = row[c] if c < len(row) else ""
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


_INLINE_MARKUP = re.compile(r"\*\*|__|`")
_NUMBERED = re.compile(r"^\d+[.)]\s+")


def _plain(text: str) -> str:
    return _INLINE_MARKUP.sub("", text).strip()


def markdown_to_docx_args(title: str, markdown: str, filename: str) -> MakeDocxArgs:
    """Best-effort Markdown -> sections: `#` sets the title, `##`/`###` start sections, `-`/`*`/
    numbered items become bullets, pipe tables become tables, other lines become paragraphs.
    Used as a safety net when an agent answered a document task in text without calling
    make_docx."""
    sections: list[DocxSection] = []
    current = DocxSection(heading="Overview", paragraphs=[], bullets=[], table_rows=[])
    paragraph: list[str] = []

    def flush_paragraph() -> None:
        if paragraph:
            current.paragraphs.append(_plain(" ".join(paragraph)))
            paragraph.clear()

    for raw in markdown.splitlines():
        line = raw.strip()
        if not line or set(line) <= {"-", "*", "_"}:  # blank line or horizontal rule
            flush_paragraph()
        elif line.startswith("# ") and not sections and not current.paragraphs:
            title = _plain(line[2:]) or title
        elif line.startswith("#"):
            flush_paragraph()
            if current.paragraphs or current.bullets or current.table_rows:
                sections.append(current)
            current = DocxSection(
                heading=_plain(line.lstrip("#")), paragraphs=[], bullets=[], table_rows=[]
            )
        elif line.startswith("|"):
            flush_paragraph()
            cells = [_plain(c) for c in line.strip("|").split("|")]
            if not all(set(c) <= {"-", ":", " "} for c in cells):  # skip |---|---| separators
                current.table_rows.append(cells)
        elif line.startswith(("- ", "* ", "+ ")) or _NUMBERED.match(line):
            flush_paragraph()
            current.bullets.append(_plain(_NUMBERED.sub("", line.lstrip("-*+ "))))
        else:
            paragraph.append(line.lstrip("> "))
    flush_paragraph()
    if current.paragraphs or current.bullets or current.table_rows or not sections:
        sections.append(current)
    return MakeDocxArgs(filename=filename, title=title, sections=sections[:MAX_SECTIONS])


async def make_docx(ctx: ToolContext, args: MakeDocxArgs) -> str:
    filename = args.filename if args.filename.lower().endswith(".docx") else f"{args.filename}.docx"
    data = render_docx(args)
    target = ctx.workspace.write_bytes(filename, data)
    relative = ctx.workspace.relative(target)
    return f"created {relative} ({len(data)} bytes, {len(args.sections)} sections)"


MAKE_DOCX = Tool(
    name="make_docx",
    description=(
        "Create a professionally formatted Word (.docx) document from a title and sections "
        "(heading, paragraphs, bullets and an optional table). Use this for every document "
        "deliverable; the file is delivered to the user as an artifact."
    ),
    args=MakeDocxArgs,
    handler=make_docx,
)
